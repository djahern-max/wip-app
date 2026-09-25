"""Source kind ``estimate_template`` (F06, D-32): the platform's own estimate
template, the one estimate source. A workbook (``.xlsx``) with the sheets
*Estimates*, *Work areas* and *Estimate costs* found by name (a *Read me* sheet or
any other is ignored), or one ``.csv`` per sheet recognised by its column header.
Columns are found by header name, trimmed and case-folded; extra columns are
ignored; a required column missing from a sheet fails the parse (the F03 sentence:
check that it is the right file).

Per estimate id the parser yields up to three ``RawItem``s, so an identical
re-export writes nothing (D-20) and a changed sheet is a new raw version of that
one entity: ``estimate`` (the Estimates row), ``work_areas`` (that id's rows on
Work areas) and ``cost_lines`` (that id's rows on Estimate costs, same-code lines
under one work area summed at parse: owner, 2026-09-25). Rows the parser cannot
load are ``RejectedItem``s with a sentence (and the code ``EST_NO_ID`` when the
row has no estimate id; its cells travel in ``detail``); the rest of the file
loads (CLAUDE.md: importers never fail a whole file for bad rows).

Money and hours are read through ``app.integrations.cells``: never a ``float`` by
our hand, quantized half up to the cent at parse.
"""

import csv
import io
from collections.abc import Iterable
from decimal import Decimal
from typing import BinaryIO

from app.integrations.base import RawItem, RejectedItem, SourceKind, register_source_kind
from app.integrations.cells import (
    CellError,
    cell_date,
    cell_hours,
    cell_int,
    cell_money,
    cell_text,
)

SOURCE = "template"
XLSX_MAGIC = b"PK\x03\x04"

SHEET_ESTIMATES = "estimates"
SHEET_WORK_AREAS = "work areas"
SHEET_COSTS = "estimate costs"
SHEET_TITLES = {
    SHEET_ESTIMATES: "Estimates",
    SHEET_WORK_AREAS: "Work areas",
    SHEET_COSTS: "Estimate costs",
}

ESTIMATE_COLUMNS = (
    "estimate_id",
    "estimator",
    "client",
    "jobsite",
    "name",
    "status",
    "price",
    "estimate_date",
)
WORK_AREA_COLUMNS = ("estimate_id", "order", "kept", "name", "price", "notes")
COST_COLUMNS = ("estimate_id", "order", "cost_code", "hours", "amount", "notes")
COLUMNS = {
    SHEET_ESTIMATES: ESTIMATE_COLUMNS,
    SHEET_WORK_AREAS: WORK_AREA_COLUMNS,
    SHEET_COSTS: COST_COLUMNS,
}
REQUIRED = {
    SHEET_ESTIMATES: ("estimate_id", "name", "status", "price"),
    SHEET_WORK_AREAS: ("estimate_id", "order", "kept", "name", "price"),
    SHEET_COSTS: ("estimate_id", "order", "cost_code", "amount"),
}
STATUS_NORM = {"pending": "pending", "sold": "sold", "lost": "lost"}
KEPT = {"Y": True, "N": False}
ID_MAX = 80


class TemplateError(ValueError):
    """The file is not the template (no template sheet, a required column missing).
    Raised from ``parse``: the batch fails with the F03 sentence."""


Row = tuple[int, list]  # (row number in the sheet, cells)


def _key(header) -> str:
    return " ".join(cell_text(header).split()).casefold()


def _sheets_from_xlsx(stream: BinaryIO) -> dict[str, list[Row]]:
    from openpyxl import load_workbook

    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        found: dict[str, list[Row]] = {}
        for ws in wb.worksheets:
            key = _key(ws.title)
            if key in COLUMNS and key not in found:
                found[key] = [(n, list(r)) for n, r in enumerate(ws.iter_rows(values_only=True), 1)]
        return found
    finally:
        wb.close()


def _sheet_from_csv(data: bytes) -> dict[str, list[Row]]:
    """One CSV is one sheet; its header says which (the most specific match wins:
    a header with ``cost_code`` is Estimate costs, with ``kept`` Work areas, with
    ``status`` Estimates)."""
    rows = [
        (n, list(r)) for n, r in enumerate(csv.reader(io.StringIO(data.decode("utf-8-sig"))), 1)
    ]
    header = next((cells for _, cells in rows if any(cell_text(c) for c in cells)), None)
    if header is None:
        raise TemplateError("empty file")
    keys = {_key(h) for h in header}
    for sheet, marker in (
        (SHEET_COSTS, "cost_code"),
        (SHEET_WORK_AREAS, "kept"),
        (SHEET_ESTIMATES, "status"),
    ):
        if marker in keys and set(REQUIRED[sheet]) <= keys:
            return {sheet: rows}
    raise TemplateError("the header names no template sheet")


def _columns(sheet: str, rows: list[Row]) -> tuple[dict[str, int], int]:
    """(column name → index, header row number). Header = first filled row."""
    for n, cells in rows:
        if any(cell_text(c) for c in cells):
            index = {}
            for i, h in enumerate(cells):
                k = _key(h)
                if k in COLUMNS[sheet] and k not in index:
                    index[k] = i
            missing = [c for c in REQUIRED[sheet] if c not in index]
            if missing:
                raise TemplateError(f"sheet {sheet!r} is missing column(s) {', '.join(missing)}")
            return index, n
    raise TemplateError(f"sheet {sheet!r} is empty")


def _reject(
    sheet: str, n: int, what: str, *, code: str | None = None, detail: dict | None = None
) -> RejectedItem:
    title = SHEET_TITLES[sheet]
    return RejectedItem(
        reason=code or "row_not_loaded",
        row_number=n,
        code=code,
        message=f'Row {n} on "{title}" was not loaded: {what}',
        detail={"sheet": title, "row": n, **(detail or {})},
    )


def _no_id(sheet: str, n: int, cells: dict[str, str]) -> RejectedItem:
    from app.domain.estimates.exceptions import sentence

    title = SHEET_TITLES[sheet]
    return RejectedItem(
        reason="EST_NO_ID",
        row_number=n,
        code="EST_NO_ID",
        message=sentence("EST_NO_ID", row=n, sheet=title),
        detail={"sheet": title, "row": n, "cells": cells},
    )


def _blank(cells: list) -> bool:
    return not any(cell_text(c) for c in cells)


def _get(cells: list, index: dict[str, int], name: str):
    i = index.get(name)
    return None if i is None or i >= len(cells) else cells[i]


def _plain_cells(cells: list, index: dict[str, int]) -> dict[str, str]:
    return {name: cell_text(_get(cells, index, name)) for name in index}


def parse_template(stream: BinaryIO) -> Iterable[RawItem | RejectedItem]:
    head = stream.read(4)
    stream.seek(0)
    sheets = _sheets_from_xlsx(stream) if head == XLSX_MAGIC else _sheet_from_csv(stream.read())
    if not sheets:
        raise TemplateError("no template sheet (Estimates, Work areas, Estimate costs) found")
    rejected: list[RejectedItem] = []
    estimates: dict[str, dict] = {}
    work_areas: dict[str, dict[int, dict]] = {}
    lines: dict[str, dict[tuple[int, str], dict]] = {}
    ids_seen_on_estimates: set[str] = set()

    def estimate_id(sheet: str, n: int, cells: list, index: dict[str, int]) -> str | None:
        eid = cell_text(_get(cells, index, "estimate_id"))
        if not eid:
            rejected.append(_no_id(sheet, n, _plain_cells(cells, index)))
            return None
        if len(eid) > ID_MAX:
            rejected.append(
                _reject(sheet, n, f"the estimate id is longer than {ID_MAX} characters.")
            )
            return None
        return eid

    if SHEET_ESTIMATES in sheets:
        sheet = SHEET_ESTIMATES
        index, header_row = _columns(sheet, sheets[sheet])
        for n, cells in sheets[sheet]:
            if n <= header_row or _blank(cells):
                continue
            eid = estimate_id(sheet, n, cells, index)
            if eid is None:
                continue
            if eid in ids_seen_on_estimates:
                rejected.append(
                    _reject(
                        sheet, n, f"{eid} appears twice on this sheet.", detail={"estimate_id": eid}
                    )
                )
                continue
            name = cell_text(_get(cells, index, "name"))
            status = cell_text(_get(cells, index, "status"))
            if not name:
                rejected.append(_reject(sheet, n, "name is required.", detail={"estimate_id": eid}))
                continue
            if not status:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "status is required (Pending, Sold or Lost).",
                        detail={"estimate_id": eid},
                    )
                )
                continue
            try:
                price = cell_money(_get(cells, index, "price"))
                estimate_date = cell_date(_get(cells, index, "estimate_date"))
            except CellError as exc:
                rejected.append(_reject(sheet, n, f"{exc}.", detail={"estimate_id": eid}))
                continue
            if price is None:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "price is required (0.00 is allowed).",
                        detail={"estimate_id": eid},
                    )
                )
                continue
            ids_seen_on_estimates.add(eid)
            estimates[eid] = {
                "estimate_id": eid,
                "estimator": cell_text(_get(cells, index, "estimator")) or None,
                "client": cell_text(_get(cells, index, "client")) or None,
                "jobsite": cell_text(_get(cells, index, "jobsite")) or None,
                "name": name,
                "status": status,
                "status_norm": STATUS_NORM.get(status.casefold()),
                "price": price,
                "estimate_date": estimate_date,
            }

    if SHEET_WORK_AREAS in sheets:
        sheet = SHEET_WORK_AREAS
        index, header_row = _columns(sheet, sheets[sheet])
        for n, cells in sheets[sheet]:
            if n <= header_row or _blank(cells):
                continue
            eid = estimate_id(sheet, n, cells, index)
            if eid is None:
                continue
            try:
                order = cell_int(_get(cells, index, "order"))
                price = cell_money(_get(cells, index, "price"))
            except CellError as exc:
                rejected.append(_reject(sheet, n, f"{exc}.", detail={"estimate_id": eid}))
                continue
            if order is None or order < 1:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "order must be a whole number from 1 up.",
                        detail={"estimate_id": eid},
                    )
                )
                continue
            kept_text = cell_text(_get(cells, index, "kept")).upper()
            if kept_text not in KEPT:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "kept must be Y or N.",
                        detail={"estimate_id": eid, "order": order},
                    )
                )
                continue
            name = cell_text(_get(cells, index, "name"))
            if not name:
                rejected.append(
                    _reject(
                        sheet, n, "name is required.", detail={"estimate_id": eid, "order": order}
                    )
                )
                continue
            if price is None:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "price is required (0.00 is allowed).",
                        detail={"estimate_id": eid, "order": order},
                    )
                )
                continue
            per_estimate = work_areas.setdefault(eid, {})
            if order in per_estimate:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        f"work area #{order} of {eid} appears twice.",
                        detail={"estimate_id": eid, "order": order},
                    )
                )
                continue
            per_estimate[order] = {
                "order": order,
                "kept": KEPT[kept_text],
                "name": name,
                "price": price,
                "notes": cell_text(_get(cells, index, "notes")) or None,
            }

    if SHEET_COSTS in sheets:
        sheet = SHEET_COSTS
        index, header_row = _columns(sheet, sheets[sheet])
        for n, cells in sheets[sheet]:
            if n <= header_row or _blank(cells):
                continue
            eid = estimate_id(sheet, n, cells, index)
            if eid is None:
                continue
            try:
                order = cell_int(_get(cells, index, "order"))
                hours = cell_hours(_get(cells, index, "hours"))
                amount = cell_money(_get(cells, index, "amount"))
            except CellError as exc:
                rejected.append(_reject(sheet, n, f"{exc}.", detail={"estimate_id": eid}))
                continue
            if order is None or order < 1:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "order must be a whole number from 1 up.",
                        detail={"estimate_id": eid},
                    )
                )
                continue
            code = cell_text(_get(cells, index, "cost_code"))
            if not code:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        "cost_code is required.",
                        detail={"estimate_id": eid, "order": order},
                    )
                )
                continue
            if amount is None:
                rejected.append(
                    _reject(
                        sheet, n, "amount is required.", detail={"estimate_id": eid, "order": order}
                    )
                )
                continue
            if eid in work_areas and order not in work_areas[eid]:
                rejected.append(
                    _reject(
                        sheet,
                        n,
                        f'work area #{order} of {eid} is not on "Work areas" in this file.',
                        detail={"estimate_id": eid, "order": order},
                    )
                )
                continue
            note = cell_text(_get(cells, index, "notes"))
            per_estimate = lines.setdefault(eid, {})
            line = per_estimate.get((order, code))
            if line is None:
                per_estimate[(order, code)] = {
                    "order": order,
                    "cost_code": code,
                    "hours": hours,
                    "amount": amount,
                    "notes": note or None,
                }
            else:
                line["amount"] = line["amount"] + amount
                if hours is not None:
                    line["hours"] = hours if line["hours"] is None else line["hours"] + hours
                if note and note not in (line["notes"] or ""):
                    line["notes"] = f"{line['notes']}; {note}" if line["notes"] else note

    yield from rejected
    for eid in sorted(set(estimates) | set(work_areas) | set(lines)):
        if eid in estimates:
            yield RawItem(entity_type="estimate", external_id=eid, payload=estimates[eid])
        if eid in work_areas:
            rows = [work_areas[eid][o] for o in sorted(work_areas[eid])]
            yield RawItem(entity_type="work_areas", external_id=eid, payload={"work_areas": rows})
        if eid in lines:
            rows = [lines[eid][k] for k in sorted(lines[eid])]
            yield RawItem(entity_type="cost_lines", external_id=eid, payload={"lines": rows})


def money_text(value: Decimal) -> str:
    return format(value, "f")


estimate_template = register_source_kind(
    SourceKind(
        name="estimate_template",
        extensions=frozenset({"xlsx", "csv"}),
        parse=parse_template,
        source=SOURCE,
        label="Estimate template",
        records_noun="estimates",
        file_noun="the estimate template",
        description=(
            "The platform's estimate template (.xlsx with the sheets Estimates, Work areas "
            "and Estimate costs, or one .csv per sheet); estimates are created and updated "
            "by estimate id."
        ),
        after_load="estimates.normalize",
        after_load_subject="estimates",
    )
)
