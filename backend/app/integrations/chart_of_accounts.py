"""Source kind ``chart_of_accounts`` (F04): a tenant's chart as a ``.csv`` or
``.xlsx`` export. One ``RawItem`` per account (``external_id`` = the account number
as text; payload ``account_no``, ``account_name``, ``ledger_type``), one
``RejectedItem`` per row whose account-number cell is empty or is not an account
number (a title row, a section heading such as "ASSETS (1000–1999)", the legend
block at the bottom of the owner's workbook). No four-digit assumption (owner
amendment B): an account number is one or more digits, optionally with dots or
hyphens between digit groups, of any length. The column-header row is skipped
silently, blank rows too.

Layout: the chart is the first sheet of a workbook (other sheets are not read), and
the account number, the account name and the ledger type are three neighbouring
columns in that order. They need not start in the first column: the owner's
workbook keeps column A empty. The starting column is the first filled cell of the
header row ("No.", "Account", "Account number", …; a row with at least two filled
cells, so a one-cell title is not the header) or, in a file without a header, of
the first row that starts with an account number. It is fixed once for the file:
after that, a row with nothing in that column is an empty account number even when
a later cell is filled. Columns are found by position, never by guessing at names.

XLSX cells are read as text. A numeric cell is turned into text through ``Decimal``
(an integer cell as its digits); nothing on this path is money, and no value ever
becomes a ``float`` by our hand. Workbooks are opened read-only, values only.
"""

import csv
import io
import re
from collections.abc import Iterable
from decimal import Decimal
from typing import BinaryIO

from app.integrations.base import RawItem, RejectedItem, SourceKind, register_source_kind

ACCOUNT_NO = re.compile(r"^\d+(?:[.\-]\d+)*$")
HEADER_WORDS = ("account", "number", "name", "type")
# Short headers are matched whole, letters only ("No." → "no"): as substrings they
# would match ordinary words such as "Notes".
HEADER_TOKENS = frozenset({"no", "num", "acct", "acctno", "acctnum"})
XLSX_MAGIC = b"PK\x03\x04"


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # openpyxl hands back a float for a numeric cell with a decimal point (an
        # Excel artefact such as 1010.0); repr is the exact shortest form of what
        # the workbook stored, and Decimal keeps it as digits from here on.
        d = Decimal(repr(value))
        return str(d.quantize(Decimal(1))) if d == d.to_integral_value() else format(d, "f")
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _is_header(cells: list[str]) -> bool:
    """``cells`` starts at the row's first filled cell."""
    first = cells[0].lower() if cells else ""
    if not first or any(c.isdigit() for c in first) or sum(1 for c in cells if c) < 2:
        return False
    letters = re.sub(r"[^a-z]", "", first)
    return first == "#" or letters in HEADER_TOKENS or any(w in first for w in HEADER_WORDS)


def _rows_from_csv(data: bytes) -> Iterable[tuple[int, list[str]]]:
    text = data.decode("utf-8-sig")
    for n, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        yield n, [c.strip() for c in row]


def _rows_from_xlsx(stream: BinaryIO) -> Iterable[tuple[int, list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            yield n, [_cell_text(v) for v in row]
    finally:
        wb.close()


def parse_chart(stream: BinaryIO) -> Iterable[RawItem | RejectedItem]:
    head = stream.read(4)
    stream.seek(0)
    rows = _rows_from_xlsx(stream) if head == XLSX_MAGIC else _rows_from_csv(stream.read())
    header_seen = False
    start: int | None = None  # index of the account-number column, once known
    for n, cells in rows:
        cells = list(cells)
        first_filled = next((i for i, c in enumerate(cells) if c), None)
        if first_filled is None:
            continue  # blank row
        if start is None:
            if _is_header(cells[first_filled:]):
                start, header_seen = first_filled, True
                continue
            if not ACCOUNT_NO.match(cells[first_filled]):
                yield RejectedItem(reason="not_an_account_number", row_number=n)
                continue
            start = first_filled
        cells = cells[start:] + ["", "", ""]
        if not header_seen and _is_header(cells):
            header_seen = True
            continue
        account_no = cells[0]
        if not account_no:
            yield RejectedItem(reason="empty_account_number", row_number=n)
            continue
        if not ACCOUNT_NO.match(account_no):
            yield RejectedItem(reason="not_an_account_number", row_number=n)
            continue
        yield RawItem(
            entity_type="account",
            external_id=account_no,
            payload={
                "account_no": account_no,
                "account_name": cells[1],
                "ledger_type": cells[2],
            },
        )


chart_of_accounts = register_source_kind(
    SourceKind(
        name="chart_of_accounts",
        extensions=frozenset({"csv", "xlsx"}),
        parse=parse_chart,
        source="chart",
        label="Chart of accounts",
        records_noun="accounts",
        file_noun="a chart of accounts export",
        description=(
            "The tenant's chart of accounts (.csv or .xlsx); accounts are updated and "
            "mappings suggested."
        ),
        after_load="config.normalize_chart",
        after_load_subject="accounts",
    )
)
