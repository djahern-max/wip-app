"""The Cost page as a file (F10; BLUEPRINT §9: values from Decimal, a tie-out tab; D-22).
XLSX with openpyxl through the F08 writer (``app.domain.billing.export.write_money``,
``MONEY_FORMAT``, the §8.7 ``LEGEND``): two tabs, the tie-out (month × account with the
five parts and the ledger total) and the unassigned list. Nothing is recomputed: the rows
are the ones the screen shows. No PDF."""

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.domain.billing.export import LEGEND, write_money
from app.domain.cost.figures import CostTieRow
from app.domain.cost.load import UnassignedView
from app.domain.cost.models import TXN_TYPE_LABELS

TIE_HEADERS: tuple[str, ...] = (
    "Month",
    "Account",
    "On jobs",
    "Owned equipment and fuel on jobs",
    "Unassigned",
    "Labor in the ledger",
    "Outside job cost",
    "Parts",
    "Ledger total",
    "Difference",
)
UNASSIGNED_HEADERS: tuple[str, ...] = (
    "Date",
    "Document",
    "Vendor",
    "Account",
    "Amount",
    "Customer row",
    "Lines",
)


@dataclass(frozen=True)
class CostExport:
    tenant_name: str
    as_of: date
    status: str
    policy_note: str | None
    tie_out: tuple[CostTieRow, ...]
    unassigned: tuple[UnassignedView, ...]
    legend: str = LEGEND

    @property
    def title(self) -> str:
        return "Cost"

    @property
    def period(self) -> str:
        return f"Figures to date as of {self.as_of.isoformat()}"


def _cell(ws, row: int, col: int, value, *, bold: bool = False, right: bool = False) -> None:
    if isinstance(value, Decimal):
        write_money(ws.cell(row=row, column=col), value)
        return
    cell = ws.cell(row=row, column=col, value=value)
    if bold:
        cell.font = Font(bold=True)
    if right:
        cell.alignment = Alignment(horizontal="right")


def _head(ws, report: CostExport) -> int:
    ws["A1"] = f"{report.title}: the ledger's cost side by month and account"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = report.tenant_name
    ws["A3"] = report.period
    ws["A4"] = f"Tie-out: {report.status}"
    row = 5
    if report.policy_note:
        ws["A5"] = report.policy_note
        row = 6
    return row + 1


def customer_words(u: UnassignedView) -> str:
    if u.doc.customer_external_id is None:
        return "No customer"
    return u.customer_name or f"QuickBooks customer {u.doc.customer_external_id}"


def cost_xlsx(report: CostExport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tie-out"
    start = _head(ws, report)
    for c, h in enumerate(TIE_HEADERS, start=1):
        _cell(ws, start, c, h, bold=True, right=c > 2)
        ws.column_dimensions[get_column_letter(c)].width = 34 if c == 2 else 18 if c > 2 else 10
    r = start + 1
    for t in report.tie_out:
        values = (
            t.month,
            t.account.label,
            t.on_jobs,
            t.memo,
            t.unassigned,
            t.labor,
            t.outside,
            t.parts,
            t.total,
            t.difference,
        )
        for c, v in enumerate(values, start=1):
            _cell(ws, r, c, v)
        r += 1
    _cell(ws, r + 1, 1, report.legend)
    ws.freeze_panes = ws.cell(row=start + 1, column=3)

    un = wb.create_sheet("Unassigned")
    start = _head(un, report)
    for c, h in enumerate(UNASSIGNED_HEADERS, start=1):
        _cell(un, start, c, h, bold=True, right=c in (5, 7))
        un.column_dimensions[get_column_letter(c)].width = 30 if c in (3, 4, 6) else 14
    r = start + 1
    for u in report.unassigned:
        values = (
            u.doc.txn_date.isoformat(),
            f"{TXN_TYPE_LABELS.get(u.doc.txn_type, u.doc.txn_type)} {u.doc.document}",
            u.vendor_name or "",
            ", ".join(u.doc.account_labels),
            u.doc.amount,
            customer_words(u),
            u.doc.lines,
        )
        for c, v in enumerate(values, start=1):
            _cell(un, r, c, v, right=c == 7)
        r += 1
    if not report.unassigned:
        _cell(un, r, 1, "Nothing unassigned: every job-cost line has a job.")
        r += 1
    _cell(un, r + 1, 1, report.legend)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
