"""The Sold Jobs Board as a file (F08; BLUEPRINT §9: "XLSX exports contain values, not
float artifacts, with a tie-out tab"; D-22; D-40). Both files are written from the same
``Decimal`` figures the screen shows; nothing is recomputed. XLSX with openpyxl (the
estimate template's library), PDF with reportlab.

The report carries its period (figures to date as of the tenant's today), the tenant
name, the tie-out status and the §8.7 legend. Money shows cents, negatives in
parentheses, zero as 0.00; a figure not shown is said in words."""

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.domain.billing.board import TieRow
from app.domain.billing.figures import words

LEGEND = "Prepared by management from the company's records. No assurance is provided."
MONEY_FORMAT = '#,##0.00;(#,##0.00);"0.00"'  # cents; negatives in parentheses; zero 0.00

COLUMNS: tuple[tuple[str, str], ...] = (
    ("job", "Job"),
    ("estimate_number", "Estimate"),
    ("customer_name", "Customer"),
    ("division_code", "Division"),
    ("revenue_method_label", "Revenue method"),
    ("status_label", "Status"),
    ("estimator", "Estimator"),
    ("revised_contract", "Revised contract"),
    ("deposit_invoiced", "Deposit invoiced"),
    ("deposit_received", "Deposit received"),
    ("billed_to_date", "Billed to date"),
    ("fuel_surcharge_billed", "Fuel surcharge billed"),
    ("collected_to_date", "Collected to date"),
    ("other_credits_applied", "Other credits applied"),
    ("open_ar", "Open A/R"),
    ("remaining_to_bill", "Remaining to bill"),
    ("days_since_activity", "Days since last activity"),
)
MONEY_KEYS = frozenset(
    {
        "revised_contract",
        "deposit_invoiced",
        "deposit_received",
        "billed_to_date",
        "fuel_surcharge_billed",
        "collected_to_date",
        "other_credits_applied",
        "open_ar",
        "remaining_to_bill",
    }
)


@dataclass(frozen=True)
class BoardRow:
    """One row as the screen shows it: text, Decimal money (None = not shown, with the
    words in ``notes``), and the day count."""

    values: dict[str, str | Decimal | int | None]
    notes: dict[str, str]  # column key → the words shown where a figure is not


@dataclass(frozen=True)
class BoardReport:
    tenant_name: str
    as_of: date
    filters: str  # "All jobs" or the filters in words
    rows: tuple[BoardRow, ...]
    totals: BoardRow
    not_on_a_job: BoardRow
    tie_out: tuple[TieRow, ...]
    tie_out_status: str
    policy_note: str | None
    legend: str = LEGEND

    @property
    def title(self) -> str:
        return "Sold Jobs Board"

    @property
    def period(self) -> str:
        return f"Figures to date as of {self.as_of.isoformat()}"


def _write_money(cell, value: Decimal) -> None:
    """A numeric cell holding exactly the Decimal's digits. openpyxl formats a number
    with ``%.16g`` (72570.85 becomes 72570.85000000001); a string marked numeric is
    written as it is, so the sheet holds the cents and nothing else."""
    cell.value = format(value.quantize(Decimal("0.01")), "f")
    cell.data_type = "n"
    cell.number_format = MONEY_FORMAT
    cell.alignment = Alignment(horizontal="right")


def _cell_text(row: BoardRow, key: str) -> str:
    value = row.values.get(key)
    if value is None:
        return row.notes.get(key, "")
    if isinstance(value, Decimal):
        return words(value)
    return str(value)


# --- XLSX ----------------------------------------------------------------------------------------


def _write_header(ws, report: BoardReport) -> None:
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = report.tenant_name
    ws["A3"] = report.period
    ws["A4"] = report.filters
    ws["A5"] = f"Tie-out: {report.tie_out_status}"
    if report.policy_note:
        ws["A6"] = report.policy_note


def _write_row(ws, r: int, row: BoardRow, *, bold: bool = False) -> None:
    for c, (key, _label) in enumerate(COLUMNS, start=1):
        value = row.values.get(key)
        cell = ws.cell(row=r, column=c)
        if value is None:
            cell.value = row.notes.get(key, "")
            if key in MONEY_KEYS or key == "days_since_activity":
                cell.alignment = Alignment(horizontal="right")
        elif isinstance(value, Decimal):
            _write_money(cell, value)
        else:
            cell.value = value
            if key == "days_since_activity":
                cell.alignment = Alignment(horizontal="right")
        if bold:
            cell.font = Font(bold=True)


def board_xlsx(report: BoardReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sold Jobs Board"
    _write_header(ws, report)
    start = 8
    for c, (_key, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=start, column=c, value=label)
        cell.font = Font(bold=True)
        if _key in MONEY_KEYS or _key == "days_since_activity":
            cell.alignment = Alignment(horizontal="right")
    r = start + 1
    for row in report.rows:
        _write_row(ws, r, row)
        r += 1
    _write_row(ws, r, report.totals, bold=True)
    _write_row(ws, r + 1, report.not_on_a_job, bold=True)
    ws.cell(row=r + 3, column=1, value=report.legend)
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    for c, (key, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = (
            34 if key == "job" else max(14, len(label) + 2)
        )

    tie = wb.create_sheet("Tie-out")
    tie["A1"] = f"Tie-out to the QuickBooks month totals: {report.tenant_name}"
    tie["A1"].font = Font(bold=True)
    tie["A2"] = report.tie_out_status
    headers = (
        "Month",
        "Jobs billed",
        "Not on a job billed",
        "Jobs + not on a job",
        "QuickBooks invoices − credit memos + sales receipts",
        "Difference",
        "Jobs collected",
        "Not on a job collected",
        "Jobs + not on a job",
        "QuickBooks payments + sales receipts",
        "Difference",
    )
    for c, h in enumerate(headers, start=1):
        cell = tie.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True)
        tie.column_dimensions[get_column_letter(c)].width = 16 if c > 1 else 10
    for i, t in enumerate(report.tie_out, start=5):
        values = (
            t.month,
            t.jobs_billed,
            t.other_billed,
            t.jobs_billed + t.other_billed,
            t.ledger_billed,
            t.billed_difference,
            t.jobs_collected,
            t.other_collected,
            t.jobs_collected + t.other_collected,
            t.ledger_collected,
            t.collected_difference,
        )
        for c, v in enumerate(values, start=1):
            if isinstance(v, Decimal):
                _write_money(tie.cell(row=i, column=c), v)
            else:
                tie.cell(row=i, column=c, value=v)
    tie.cell(row=6 + len(report.tie_out), column=1, value=report.legend)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- PDF -----------------------------------------------------------------------------------------


def _pdf_table(data: list[list[str]], right: set[int], *, bold_last: int = 0) -> Table:
    table = Table(data, repeatRows=1)
    style = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEADING", (0, 0), (-1, -1), 8),
    ]
    for c in right:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    if bold_last:
        style.append(("FONTNAME", (0, -bold_last), (-1, -1), "Helvetica-Bold"))
        style.append(("LINEABOVE", (0, -bold_last), (-1, -bold_last), 0.5, colors.black))
    table.setStyle(TableStyle(style))
    return table


def board_pdf(report: BoardReport) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(letter),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"{report.title} · {report.tenant_name}",
        pageCompression=0,  # a board is small; the text stays readable to a test and a grep
    )
    styles = getSampleStyleSheet()
    small = styles["BodyText"]
    small.fontSize = 8
    small.leading = 10
    story = [
        Paragraph(report.title, styles["Heading2"]),
        Paragraph(report.tenant_name, small),
        Paragraph(report.period, small),
        Paragraph(report.filters, small),
        Paragraph(f"Tie-out: {report.tie_out_status}", small),
    ]
    if report.policy_note:
        story.append(Paragraph(report.policy_note, small))
    story.append(Spacer(1, 6))
    header = [label for _key, label in COLUMNS]
    body = [[_cell_text(r, key) for key, _ in COLUMNS] for r in report.rows]
    body.append([_cell_text(report.totals, key) for key, _ in COLUMNS])
    body.append([_cell_text(report.not_on_a_job, key) for key, _ in COLUMNS])
    right = {
        i for i, (key, _) in enumerate(COLUMNS) if key in MONEY_KEYS or key == "days_since_activity"
    }
    story.append(_pdf_table([header, *body], right, bold_last=2))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Tie-out to the QuickBooks month totals", styles["Heading4"]))
    tie_header = [
        "Month",
        "Jobs billed",
        "Not on a job",
        "QuickBooks billed",
        "Difference",
        "Jobs collected",
        "Not on a job",
        "QuickBooks collected",
        "Difference",
    ]
    tie_body = [
        [
            t.month,
            words(t.jobs_billed),
            words(t.other_billed),
            words(t.ledger_billed),
            words(t.billed_difference),
            words(t.jobs_collected),
            words(t.other_collected),
            words(t.ledger_collected),
            words(t.collected_difference),
        ]
        for t in report.tie_out
    ]
    story.append(_pdf_table([tie_header, *tie_body], set(range(1, 9))))
    story.append(Spacer(1, 8))
    story.append(Paragraph(report.legend, small))
    doc.build(story)
    return out.getvalue()
