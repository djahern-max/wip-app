"""F10 (BLUEPRINT §8.5 check 2; D-51): the constructed cost fixture through ``apply_raw``
on a tenant configured as Rye Beach (the F07 helpers), with the 67 Elm Street job linked to
project 201 as Excavation, the Turley job to 202 as Landscaping and a pool job to 205.
The README of ``tests/fixtures/rye_beach/cost`` is the oracle; nothing here is derived
from the code under test. The F08.2 pattern: for every month and account, the parts the
pure figures give equal the sum of the lines held and the raw re-sum."""

import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select

from app.core.db import tenant_session
from app.domain.cost import figures, load, report
from app.domain.cost.models import LedgerLine
from app.domain.cost.run import COST_RUN_KIND
from app.domain.exceptions.models import ReviewException
from app.domain.exceptions.run import refresh
from app.worker.models import Task
from tests.conftest import Seed
from tests.cost_helpers import (
    apply_all,
    apply_version,
    fixture,
    job_oracle,
    ledger_oracle,
    lines,
    record,
    store_fixture,
)
from tests.job_helpers import ELM_ID, TURLEY_ID, Tenant, make_tenant
from tests.test_qbo_tasks import _drain

D = Decimal
ELM_CUSTOMER, TURLEY_CUSTOMER, POOL_CUSTOMER = "201", "202", "205"
UNASSIGNED = "Bill A-1042, 2,000.00 to 5240 Subcontractors - EX, is coded to no customer."
MISMATCH = (
    "Bill A-1052 posts 4,000.00 to 5240 Subcontractors - EX (Excavation) on a Landscaping job."
)


@dataclass
class CostTenant(Tenant):
    """The F07 tenant with the three jobs the cost fixture is keyed to."""

    elm: dict = field(default_factory=dict)
    turley: dict = field(default_factory=dict)
    pool: dict = field(default_factory=dict)


@pytest.fixture
def t(
    seed: Seed, rw_engine: Engine, login_as: Callable[..., TestClient], fresh_tenant
) -> CostTenant:
    tenant = CostTenant(**vars(make_tenant(seed, rw_engine, login_as, fresh_tenant)))
    tenant.elm = tenant.new_job(ELM_ID, division="EX")
    tenant.link(tenant.elm["id"], "elm", in_progress=True)
    tenant.turley = tenant.new_job(TURLEY_ID, division="LS")
    tenant.link(tenant.turley["id"], "turley", in_progress=True)
    tenant.pool = tenant.send(
        "POST",
        "/api/jobs/pool",
        {"name": "Pool - Hydroseed", "division_id": str(tenant.divisions["LS"])},
        201,
    )
    tenant.link(tenant.pool["id"], "pool")
    _drain(rw_engine, fresh_tenant)  # the links' cost runs find no lines yet
    store_fixture(rw_engine, fresh_tenant)
    return tenant


def _confirm_maps(t: Tenant) -> None:
    """The chart with the Rye Beach rules suggested; confirm every suggestion (F04)."""
    from tests.config_helpers import CHART_CSV, upload_chart

    upload_chart(t.client, CHART_CSV.read_bytes())
    from tests.config_helpers import run_until_quiet

    run_until_quiet(t.engine)
    t.send("POST", "/api/config/accounts/confirm-all")


def _by_doc(rows: list[LedgerLine]) -> dict[tuple[str, str], list[LedgerLine]]:
    out: dict[tuple[str, str], list[LedgerLine]] = {}
    for r in rows:
        out.setdefault((r.txn_type, r.txn_external_id), []).append(r)
    return out


def _tie(t: Tenant) -> tuple[list[figures.CostTieRow], dict, dict]:
    with tenant_session(t.engine, t.id) as s:
        r = load.cost_report(s, t.id)
        return r.tie_out, report.ledger_sums(s, t.id), report.raw_sums(s, t.id)


def _check_tie(t: Tenant) -> list[figures.CostTieRow]:
    """The F08.2 pattern: parts = Σ lines held = the raw re-sum, every month and account."""
    rows, ledger, raw = _tie(t)
    for r in rows:
        assert r.balanced, (r.month, r.account.label)
        assert (
            r.parts == ledger[r.month][r.account.key] == raw[r.month].get(r.account.key, D("0.00"))
        ), (
            r.month,
            r.account.label,
        )
    assert {(m, a) for m, cells in ledger.items() for a in cells} == {
        (r.month, r.account.key) for r in rows
    }
    return rows


def _cost(t: Tenant, job: dict) -> dict:
    return t.job(job["id"])["cost"]


def _slot(cost: dict, slot: str) -> str | None:
    return next(c["cost_to_date"] for c in cost["categories"] if c["slot"] == slot)


# --- the normalizer against the oracle  ----------------------------------------------------------


def test_normalizing_the_fixture_writes_exactly_the_oracles_lines(t: Tenant) -> None:
    _confirm_maps(t)
    outcomes = apply_all(t.engine, t.id)
    assert {o.result for o in outcomes} == {"applied"} and len(outcomes) == 11
    rows = lines(t.engine, t.id)
    assert len(rows) == 20 and all(r.deleted_at is None for r in rows)
    docs = _by_doc(rows)
    elm_job, turley_job, pool_job = (uuid.UUID(j["id"]) for j in (t.elm, t.turley, t.pool))
    # Signs as specified.
    (vc,) = docs[("VendorCredit", "7701")]
    assert vc.amount == D("-150.00") and vc.job_id == elm_job and vc.vendor_external_id == "7001"
    (check,) = docs[("Purchase", "5002")]
    assert check.amount == D("-400.00")
    je = sorted(docs[("JournalEntry", "9001")], key=lambda r: r.line_no)
    assert [(r.account_external_id, r.amount, r.job_id) for r in je] == [
        ("301", D("12500.00"), None),
        ("302", D("9800.00"), None),
    ]
    (dep,) = docs[("Deposit", "3301")]
    assert (
        dep.amount == D("-275.00") and dep.job_id == elm_job and dep.customer_external_id == "201"
    )
    # The item-based line carries the item's expense account; the one without has none.
    stone, service = sorted(docs[("Bill", "98894")], key=lambda r: r.line_no)
    assert stone.account_external_id == "303" and stone.gl_account_id is not None
    assert (
        service.account_external_id is None
        and service.gl_account_id is None
        and service.job_id == elm_job
    )
    # No CustomerRef: no job and no customer.
    (loose,) = docs[("Bill", "98893")]
    assert loose.job_id is None and loose.customer_external_id is None
    # The pool bill's seven lines carry the pool project and its job.
    pool_lines = docs[("Bill", "98890")]
    assert len(pool_lines) == 7 and all(
        r.customer_external_id == "205" and r.job_id == pool_job for r in pool_lines
    )
    # The Landscaping job's line.
    (wall,) = docs[("Bill", "98908")]
    assert wall.job_id == turley_job and wall.amount == D("4000.00")
    # Every stored line is on an expense-side account or has none; the memo is the document's.
    assert all(r.account_external_id in ("301", "302", "303", "304", "305", None) for r in rows)
    assert docs[("Bill", "98892")][0].memo == "Ramp bill; receipt in Ramp"
    # The gl_account rows were attached by number.
    with tenant_session(t.engine, t.id) as s:
        held = {
            r.account_external_id
            for r in s.execute(select(LedgerLine)).scalars()
            if r.gl_account_id
        }
    assert held == {"301", "302", "303", "304", "305"}


def test_applying_the_same_versions_twice_changes_nothing(t: Tenant) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    before = [
        (r.id, r.amount, r.job_id, r.deleted_at, r.raw_record_id) for r in lines(t.engine, t.id)
    ]
    outcomes = apply_all(t.engine, t.id)
    assert all(o.reason == "unchanged" for o in outcomes)
    assert [
        (r.id, r.amount, r.job_id, r.deleted_at, r.raw_record_id) for r in lines(t.engine, t.id)
    ] == before


def test_a_delete_stub_marks_every_line_and_a_shorter_version_drops_a_line(t: Tenant) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    elm_before = _cost(t, t.elm)
    assert elm_before["cost_to_date"] == "2361.25"
    outcome = apply_version(
        t.engine, t.id, "Bill", {**fixture("cdc_delete_stub"), "Id": "98895"}, deleted=True
    )
    assert outcome.result == "deleted"
    docs = _by_doc(lines(t.engine, t.id))
    assert all(r.deleted_at is not None for r in docs[("Bill", "98895")])
    after = _cost(t, t.elm)
    assert after["cost_to_date"] == "1861.25" and _slot(after, "40") == "425.00"
    rows = _check_tie(t)
    oct_sub = next(r for r in rows if r.month == "2026-10" and r.account.account_no == "5240")
    assert oct_sub.total == D("3725.00")
    # Bill 98892 returns with one line fewer: the Materials line goes.
    assert (
        apply_version(t.engine, t.id, "Bill", record("Bill_98892_one_line", "98892")).reason
        == "changed"
    )
    docs = _by_doc(lines(t.engine, t.id))
    assert [r.line_external_id for r in docs[("Bill", "98892")]] == ["1"]
    assert _slot(_cost(t, t.elm), "30") == "955.75"
    _check_tie(t)
    # A voided document: every line marked, no figure includes it.
    voided = record("Bill", "98908")
    voided["PrivateNote"] = "Voided"
    voided["TotalAmt"] = 0
    for ln in voided["Line"]:
        ln["Amount"] = 0
    assert apply_version(t.engine, t.id, "Bill", voided).reason == "changed"
    assert all(r.deleted_at is not None for r in _by_doc(lines(t.engine, t.id))[("Bill", "98908")])
    assert _slot(_cost(t, t.turley), "40") == "0.00"


def test_cost_to_date_on_elm_equals_the_oracle_and_labor_is_only_in_the_ledger(t: Tenant) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    expected = job_oracle("ELM")
    cost = _cost(t, t.elm)
    assert cost["cost_to_date"] == str(expected["in_basis"]) == "2361.25"
    assert cost["memo_cost_to_date"] == str(expected["memo"]) == "0.00"
    by_slot = {c["slot"]: c for c in cost["categories"]}
    assert by_slot["40"]["cost_to_date"] == str(expected["40"]) == "925.00"
    assert by_slot["30"]["cost_to_date"] == str(expected["30"]) == "1436.25"
    assert by_slot["10"]["cost_to_date"] is None and "hours and rates" in by_slot["10"]["note"]
    assert by_slot["20"]["cost_to_date"] is None
    assert all(c["in_basis"] is not None for c in cost["categories"])
    assert (
        cost["lines"] == 8
    )  # 98892 ×2, VC 7701, 5001, 5002, 3301, 98894 stone, 98895; not the 90.00 line
    # Estimated by category is the F06.1 figure: the estimate's own detail agrees.
    detail = t.get(f"/api/estimates/{t.estimate_id(ELM_ID)}")
    est_by_slot = {c["slot"]: c for c in detail["totals"]["by_category"]}
    assert by_slot["40"]["estimated"] == est_by_slot["40"]["amount"]
    assert by_slot["30"]["estimated"] == est_by_slot["30"]["amount"]
    assert by_slot["20"]["estimated"] == est_by_slot["20"]["amount_with_burden"]
    assert D(by_slot["40"]["difference"]) == D(by_slot["40"]["estimated"]) - D("925.00")
    assert cost["estimated_in_basis"] == t.job(t.elm["id"])["eac_in_basis"]
    # The pool job and the Turley job.
    assert _slot(_cost(t, t.pool), "35") == str(job_oracle("POOL")["35"]) == "8120.00"
    assert _slot(_cost(t, t.turley), "40") == str(job_oracle("TURLEY")["40"]) == "4000.00"
    # The job page's fact: cost to date in the WIP basis, and the row.
    page = t.job(t.elm["id"])
    assert page["cost"]["cost_to_date_note"] is None
    # Labor: the two JE lines are in no job and only under labor in the ledger.
    with tenant_session(t.engine, t.id) as s:
        r = load.cost_report(s, t.id)
    assert [(row.month, row.division_code, row.amount) for row in r.labor] == [
        ("2026-09", "EX", D("9800.00")),
        ("2026-09", "LS", D("12500.00")),
    ]
    assert not any(u.doc.txn_type == "JournalEntry" for u in r.unassigned)


def test_for_every_month_and_account_the_parts_equal_the_lines_and_the_raw_re_sum(
    t: Tenant,
) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    rows = _check_tie(t)
    expected = ledger_oracle()
    cells = {(r.month, r.account.account_no or ""): r for r in rows}
    assert set(cells) == set(expected)
    for key, want in expected.items():
        r = cells[key]
        got = {
            "on_jobs": r.on_jobs,
            "memo": r.memo,
            "unassigned": r.unassigned,
            "labor": r.labor,
            "outside": r.outside,
            "total": r.total,
        }
        assert got == want, key
        # The brief's four parts equal the total on every mapped job-cost account.
        if r.account.job_cost or r.account.labor:
            assert r.on_jobs + r.unassigned + r.labor + r.memo == r.total, key
    assert cells[("2026-10", "")].account.label == "No account"
    # The Cost page and its export carry the same rows.
    page = t.get("/api/cost")
    assert (
        page["balanced"] and page["months"] == 2 and page["accounts"] == 6 and page["skipped"] == 0
    )
    assert page["status"] == "Ties to the cent for 2 months and 6 accounts; every payload read."
    assert page["lines_held"] == 20
    assert [(r["month"], r["account_label"], r["total"]) for r in page["tie_out"]][:2] == [
        ("2026-09", "5110 Gross Payroll - LS", "12500.00"),
        ("2026-09", "5135 Supplies - LS", "8120.00"),
    ]
    assert [(u["document"], u["amount"], u["customer_label"]) for u in page["unassigned"]] == [
        ("A-1042", "2000.00", "No customer")
    ]
    assert page["unassigned"][0]["vendor"] == "Vendor A Concrete"
    r = t.client.get("/api/cost/export.xlsx")
    assert r.status_code == 200
    import io
    import zipfile

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tie-out", "Unassigned"]
    tie = wb["Tie-out"]
    header_row = next(i for i in range(1, 10) if tie.cell(row=i, column=1).value == "Month")
    assert [tie.cell(row=header_row + 1, column=c).value for c in (1, 2)] == [
        "2026-09",
        "5110 Gross Payroll - LS",
    ]
    un = wb["Unassigned"]
    header_row = next(i for i in range(1, 10) if un.cell(row=i, column=1).value == "Date")
    assert [un.cell(row=header_row + 1, column=c).value for c in (2, 6)] == [
        "Bill A-1042",
        "No customer",
    ]
    # The sheets hold the Decimal's digits, never a float's (the F08 writer): the cells are
    # written as the strings "12500.00" and "2000.00" marked numeric.
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        sheets = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/"))
    assert "<v>12500.00</v>" in sheets and "<v>2000.00</v>" in sheets and "<v>8120.00</v>" in sheets
    assert "12500.0<" not in sheets and "0000001" not in sheets


# --- links fill job_id on the next run; the run is enqueued by the link  -------------------------


def test_linking_a_project_after_the_lines_fills_job_id_on_the_next_run_and_unlinking_clears_it(
    t: Tenant,
) -> None:
    _confirm_maps(t)
    # Unlink the pool first, so its lines are unassigned when they are normalized.
    pool_detail = t.job(t.pool["id"])
    alias = next(a for a in pool_detail["aliases"] if a["system"] == "qbo_customer")
    t.send("DELETE", f"/api/jobs/{t.pool['id']}/aliases/{alias['id']}")
    _drain(t.engine, t.id)
    apply_all(t.engine, t.id)
    docs = _by_doc(lines(t.engine, t.id))
    assert all(
        r.job_id is None and r.customer_external_id == "205" for r in docs[("Bill", "98890")]
    )
    page = t.get("/api/cost")
    assert any(
        u["document"] == "S-3919" and u["customer_label"] == "Pool - Hydroseed"
        for u in page["unassigned"]
    )
    # The link queues one cost run (job ids only); the worker runs it; the lines take the job.
    with tenant_session(t.engine, t.id) as s:
        before = s.execute(select(Task).where(Task.kind == COST_RUN_KIND)).scalars().all()
        n_before = len(before)
    t.link(t.pool["id"], "pool")
    with tenant_session(t.engine, t.id) as s:
        queued = [
            x for x in s.execute(select(Task).where(Task.kind == COST_RUN_KIND)).scalars().all()
        ]
        assert len(queued) == n_before + 1 and queued[-1].payload == {"full": False}
    _drain(t.engine, t.id)
    pool_job = uuid.UUID(t.pool["id"])
    docs = _by_doc(lines(t.engine, t.id))
    assert all(r.job_id == pool_job for r in docs[("Bill", "98890")])
    assert _slot(_cost(t, t.pool), "35") == "8120.00"
    assert not any(u["document"] == "S-3919" for u in t.get("/api/cost")["unassigned"])
    _check_tie(t)
    # Unlinking clears it on the next run.
    pool_detail = t.job(t.pool["id"])
    alias = next(a for a in pool_detail["aliases"] if a["system"] == "qbo_customer")
    t.send("DELETE", f"/api/jobs/{t.pool['id']}/aliases/{alias['id']}")
    _drain(t.engine, t.id)
    docs = _by_doc(lines(t.engine, t.id))
    assert all(r.job_id is None for r in docs[("Bill", "98890")])
    _check_tie(t)


# --- the two exceptions  -------------------------------------------------------------------------


def _run(t: Tenant):
    with tenant_session(t.engine, t.id) as s:
        return refresh(s, t.id)


def _exceptions(t: Tenant, code: str) -> list[dict]:
    return [e for e in t.get("/api/exceptions")["exceptions"] if e["code"] == code]


def test_cost_unassigned_and_division_mismatch_raise_once_resolve_and_dismiss_with_a_note(
    t: Tenant,
) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    _run(t)
    unassigned = _exceptions(t, "COST_UNASSIGNED")
    assert [e["message"] for e in unassigned] == [UNASSIGNED]
    assert unassigned[0]["subject"] == {"type": "tenant", "id": str(t.id), "label": "No customer"}
    assert unassigned[0]["severity"] == "warn" and unassigned[0]["may_dismiss"] is True
    mismatch = _exceptions(t, "COST_DIVISION_MISMATCH")
    assert [e["message"] for e in mismatch] == [MISMATCH]
    assert (
        mismatch[0]["subject"]["type"] == "job" and mismatch[0]["subject"]["id"] == t.turley["id"]
    )
    # On the job page too (one computation per subject), and on the board.
    assert MISMATCH in [i["message"] for i in t.job(t.turley["id"])["attention"]]
    board = next(j for j in t.get("/api/jobs")["jobs"] if j["id"] == t.turley["id"])
    assert MISMATCH in [i["message"] for i in board["attention"]]
    # Neither fires on a labor-slot line: the journal entry raises nothing.
    with tenant_session(t.engine, t.id) as s:
        codes = [
            (r.code, r.detail.get("txn_type")) for r in s.execute(select(ReviewException)).scalars()
        ]
    assert ("COST_UNASSIGNED", "JournalEntry") not in codes
    assert all(not (c.startswith("COST_") and txn == "JournalEntry") for c, txn in codes)
    # A second run writes nothing.
    assert _run(t).writes == 0
    # Dismiss the mismatch with a note: one audit row.
    audits_before = t.audit_rows()
    out = t.send(
        "POST",
        f"/api/exceptions/{mismatch[0]['id']}/dismiss",
        {"note": "Shared footing; owner agreed"},
    )
    assert out["status"] == "dismissed"
    assert [a.action for a in t.audit_actions(audits_before)] == ["exception_dismissed"]
    assert MISMATCH not in [i["message"] for i in t.job(t.turley["id"])["attention"]]
    assert MISMATCH in [d["message"] for d in t.job(t.turley["id"])["dismissed"]]
    # A person's edit in QuickBooks gives the line its project: the unassigned one resolves.
    apply_version(t.engine, t.id, "Bill", record("Bill_98893_assigned", "98893"))
    assert _slot(_cost(t, t.elm), "40") == "2925.00"
    result = _run(t)
    assert result.resolved == 1
    assert _exceptions(t, "COST_UNASSIGNED")[0]["status"] == "resolved"
    _check_tie(t)


# --- the roles and the words  --------------------------------------------------------------------


def test_the_cost_page_is_refused_to_client_pm_and_viewer_and_read_by_the_firm_and_client_admin(
    t: Tenant, seed: Seed, login_as, owner_engine: Engine
) -> None:
    from app.tenancy.models import Membership, Role

    _confirm_maps(t)
    apply_all(t.engine, t.id)
    added = []
    try:
        for key, role, status in (
            ("client_pm", Role.client_pm, 403),
            ("client_viewer", Role.client_viewer, 403),
            ("client_admin", Role.client_admin, 200),
            ("recover_me", None, 200),  # firm_staff, an entry row already in the fresh tenant
        ):
            if role is not None:
                with tenant_session(owner_engine, t.id) as s:
                    s.add(Membership(tenant_id=t.id, user_id=seed.users[key].id, role=role))
                added.append(seed.users[key].id)
            c = login_as(key, tenant=t.id)
            assert c.get("/api/cost").status_code == status, key
            assert c.get("/api/cost/export.xlsx").status_code == status, key
            # The job page is read by every role, with the cost block.
            page = c.get(f"/api/jobs/{t.elm['id']}")
            assert page.status_code == 200 and page.json()["cost"]["cost_to_date"] == "2361.25"
    finally:
        with tenant_session(owner_engine, t.id) as s:
            for m in s.execute(select(Membership).where(Membership.user_id.in_(added))).scalars():
                s.delete(m)
    # Money reaches the API as strings with two decimals everywhere on the page.
    page = t.get("/api/cost")
    for row in page["tie_out"]:
        for key in (
            "on_jobs",
            "memo_on_jobs",
            "unassigned",
            "labor_in_ledger",
            "outside_job_cost",
            "total",
        ):
            assert (
                isinstance(row[key], str)
                and row[key].count(".") == 1
                and len(row[key].split(".")[1]) == 2
            )


def test_an_undecided_basis_shows_not_computed_and_still_ties(t: Tenant, seed: Seed) -> None:
    from app.domain.config.models import TenantPolicy

    _confirm_maps(t)
    apply_all(t.engine, t.id)
    with tenant_session(t.engine, t.id) as s:
        row = s.execute(select(TenantPolicy).where(TenantPolicy.key == "wip_basis")).scalar_one()
        s.delete(row)
    cost = _cost(t, t.elm)
    assert cost["cost_to_date"] is None and cost["cost_to_date_note"].startswith("Not computed")
    assert all(c["in_basis"] is None for c in cost["categories"])
    page = t.get("/api/cost")
    assert page["policy_note"] and "WIP basis" in page["policy_note"]
    assert page["balanced"]
    _check_tie(t)


def test_connections_counts_the_cost_entities_held_and_the_lines(t: Tenant) -> None:
    _confirm_maps(t)
    apply_all(t.engine, t.id)
    with tenant_session(t.engine, t.id) as s:
        skipped = report.skipped_counts(s, t.id)
        held = load.lines_held(s)
    assert skipped == {"Bill": 0, "VendorCredit": 0, "Purchase": 0, "JournalEntry": 0, "Deposit": 0}
    assert held == 20
    # An unreadable payload (three decimals) is counted, never dropped silently.
    broken = record("Bill", "98893")
    broken["Line"][0]["Amount"] = D("1.005")
    outcome = apply_version(t.engine, t.id, "Bill", broken)
    assert outcome.result == "skipped" and outcome.reason == "line_amount_more_than_two_decimals"
    with tenant_session(t.engine, t.id) as s:
        assert report.skipped_counts(s, t.id)["Bill"] == 1
    assert t.get("/api/cost")["skipped"] == 1
    assert t.get("/api/cost")["status"].endswith("1 payload could not be read.")
