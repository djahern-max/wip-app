"""F08 exports: the XLSX holds the screen's values to the cent as numbers written from
Decimal (no float artifacts), with a tie-out tab that balances; the PDF (D-40) carries
the same rows and totals; both follow the filters and the role matrix of the jobs read."""

import io
import re
import zipfile
from collections.abc import Callable
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine

from app.domain.billing import export
from tests.billing_helpers import (
    DEPOSIT_ITEM,
    FUEL_ITEM,
    WORK_ITEM,
    apply_payloads,
    billing_policy,
    document_payload,
    line,
    payment_payload,
)
from tests.conftest import Seed
from tests.job_helpers import ELM_ID, TURLEY_ID, Tenant, make_tenant

D = Decimal
ELM_CUSTOMER = "201"


@pytest.fixture
def t(seed: Seed, rw_engine: Engine, login_as: Callable[..., TestClient], fresh_tenant) -> Tenant:
    tenant = make_tenant(seed, rw_engine, login_as, fresh_tenant)
    billing_policy(rw_engine, seed, fresh_tenant)
    job = tenant.new_job(ELM_ID)
    tenant.link(job["id"], "elm", in_progress=True)
    tenant.send("POST", f"/api/jobs/{job['id']}/work-areas/kinds/confirm-suggested")
    tenant.new_job(TURLEY_ID)  # sold, unlinked: backlog, no money
    apply_payloads(
        rw_engine,
        fresh_tenant,
        [
            (
                "Invoice",
                document_payload(
                    "6001",
                    customer=ELM_CUSTOMER,
                    date="2026-06-30",
                    doc_number="EST6115758_DEP",
                    lines=[line("149800.00", DEPOSIT_ITEM)],
                    balance="0",
                ),
            ),
            (
                "Invoice",
                document_payload(
                    "6002",
                    customer=ELM_CUSTOMER,
                    date="2026-09-15",
                    doc_number="EST6115758_PMT3",
                    lines=[line("70290.81", WORK_ITEM), line("3514.54", FUEL_ITEM)],
                ),
            ),
            (
                "CreditMemo",
                document_payload(
                    "6003",
                    customer=ELM_CUSTOMER,
                    date="2026-09-20",
                    lines=[line("1234.50", WORK_ITEM)],
                ),
            ),
            (
                "Payment",
                payment_payload(
                    "8001",
                    customer=ELM_CUSTOMER,
                    date="2026-07-01",
                    total="149800.00",
                    applied=[("149800.00", "Invoice", "6001")],
                ),
            ),
            (
                "Invoice",
                document_payload(
                    "6004", customer="206", date="2026-09-21", lines=[line("77.00", WORK_ITEM)]
                ),
            ),
        ],
    )
    return tenant


COLUMN_KEYS = {label: key for key, label in export.COLUMNS}
MONEY = {
    "Revised contract": "revised_contract",
    "Deposit invoiced": "deposit_invoiced",
    "Deposit received": "deposit_received",
    "Billed to date": "billed_to_date",
    "Fuel surcharge billed": "fuel_surcharge_billed",
    "Collected to date": "collected_to_date",
    "Other credits applied": "other_credits_applied",
    "Open A/R": "open_ar",
    "Remaining to bill": "remaining_to_bill",
}


def _sheet_rows(ws) -> tuple[list[str], list[list]]:
    header_row = next(i for i in range(1, 20) if ws.cell(row=i, column=1).value == "Job")
    header = [ws.cell(row=header_row, column=c).value for c in range(1, len(export.COLUMNS) + 1)]
    rows = []
    r = header_row + 1
    while ws.cell(row=r, column=1).value not in (None, ""):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, len(export.COLUMNS) + 1)])
        r += 1
    return header, rows


def test_xlsx_cells_equal_the_screen_to_the_cent_and_the_tie_out_tab_balances(t: Tenant) -> None:
    screen = t.get("/api/jobs")
    screen["tie_out"] = t.get("/api/jobs/tie-out")  # F08.2: its own request; the tab stays
    r = t.client.get("/api/jobs/export.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="sold-jobs-board-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Sold Jobs Board", "Tie-out"]
    ws = wb["Sold Jobs Board"]
    assert ws["A1"].value == "Sold Jobs Board" and ws["A2"].value == screen["tenant_name"]
    assert ws["A3"].value == f"Figures to date as of {screen['as_of']}"
    assert ws["A4"].value == "All jobs"
    assert ws["A5"].value == f"Tie-out: {screen['tie_out']['status']}"
    header, rows = _sheet_rows(ws)
    assert header == [label for _k, label in export.COLUMNS]
    assert [row[0] for row in rows] == [
        *(j["name"] for j in screen["jobs"]),
        "Total",
        "Not on a job",
    ]
    for row, job in zip(rows, screen["jobs"], strict=False):
        cells = dict(zip(header, row, strict=True))
        assert (
            cells["Estimate"] == job["estimate_number"] and cells["Estimator"] == job["estimator"]
        )
        for label, key in MONEY.items():
            api = job[key] if key == "revised_contract" else job["billing"][key]
            if api is None:
                assert isinstance(cells[label], str), (label, cells[label])  # words, never blank
            else:
                # openpyxl reads a number back as float; the sheet XML holds the digits
                # written from Decimal (checked below), so repr round-trips to the cents.
                assert D(repr(cells[label])) == D(api), (job["name"], label)
        days = job["billing"]["days_since_activity"]
        assert cells["Days since last activity"] == (days if days is not None else "No activity")
    total_cells = dict(zip(header, rows[-2], strict=True))
    other_cells = dict(zip(header, rows[-1], strict=True))
    for label, key in MONEY.items():
        if key == "revised_contract":
            continue
        for cells, api in (
            (total_cells, screen["totals"][key]),
            (other_cells, screen["not_on_a_job"][key]),
        ):
            if api is None:
                assert cells[label] in ("", None) or isinstance(cells[label], str)
            else:
                assert D(repr(cells[label])) == D(api), (label, api)
    assert D(repr(total_cells["Billed to date"])) == D("218856.31")
    assert D(repr(other_cells["Billed to date"])) == D("1077.00")  # 77.00 + the F07 seed row
    assert ws.cell(row=ws.max_row, column=1).value == export.LEGEND
    # Every money cell carries the cents format with negatives in parentheses.
    money_columns = {header.index(label) + 1 for label in MONEY}
    for row_cells in ws.iter_rows(min_row=9, max_row=ws.max_row):
        for cell in row_cells:
            if cell.column in money_columns and isinstance(cell.value, int | float):
                assert cell.number_format == export.MONEY_FORMAT
    # No float artifacts: the sheet XML holds every number with at most two decimals and
    # no exponent, exactly as the Decimal printed it.
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    numbers = re.findall(r"<v>([-0-9.Ee+]+)</v>", xml)
    odd = [n for n in numbers if not re.fullmatch(r"-?\d+(\.\d{1,2})?", n)]
    assert numbers and odd == [], odd
    assert (
        "218856.31" in numbers and "72570.85" in numbers
    )  # the cents; 72570.85 is the open A/R sum

    tie = wb["Tie-out"]
    assert tie["A2"].value == screen["tie_out"]["status"]
    months = []
    r_ = 5
    while tie.cell(row=r_, column=1).value not in (None, "", export.LEGEND):
        values = [tie.cell(row=r_, column=c).value for c in range(1, 12)]
        months.append(values[0])
        cents = [D(repr(v)) if isinstance(v, int | float) else v for v in values]
        assert cents[1] + cents[2] == cents[3] == cents[4] and cents[5] == D("0")
        assert cents[6] + cents[7] == cents[8] == cents[9] and cents[10] == D("0")
        r_ += 1
    assert months == ["2026-06", "2026-07", "2026-08", "2026-09"]  # 2026-08: the F07 seed row
    assert screen["tie_out"]["balanced"] and screen["tie_out"]["months"] == 4


def test_the_exports_follow_the_filters_and_the_pdf_carries_the_same_rows(t: Tenant) -> None:
    screen = t.get("/api/jobs?status=sold")
    screen["tie_out"] = t.get("/api/jobs/tie-out")
    (turley,) = [j["name"] for j in screen["jobs"]]
    r = t.client.get("/api/jobs/export.xlsx?status=sold")
    ws = load_workbook(io.BytesIO(r.content))["Sold Jobs Board"]
    header, rows = _sheet_rows(ws)
    assert [row[0] for row in rows] == [turley, "Total", "Not on a job"]
    assert ws["A4"].value == "Jobs with status Sold"
    # The tie-out is over every job whatever the filter.
    assert (
        load_workbook(io.BytesIO(r.content))["Tie-out"]["A2"].value == screen["tie_out"]["status"]
    )

    r = t.client.get("/api/jobs/export.pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert (
        r.content.startswith(b"%PDF")
        and 'filename="sold-jobs-board-' in r.headers["content-disposition"]
    )
    text = r.content.decode("latin-1")
    full = t.get("/api/jobs")
    for job in full["jobs"]:
        assert job["name"] in text
    for needle in (
        "Sold Jobs Board",
        full["tenant_name"],
        "Total",
        "Not on a job",
        "218,856.31",
        "1,077.00",
        export.LEGEND[:40],
    ):
        assert needle in text, needle
    assert f"Figures to date as of {full['as_of']}" in text


def test_the_pdf_and_the_xlsx_are_built_from_the_same_report() -> None:
    from datetime import date

    from app.domain.billing.board import TieRow

    row = export.BoardRow(
        values={
            "job": "A job",
            "estimate_number": "EST1",
            "customer_name": "C",
            "division_code": "EX",
            "revenue_method_label": "Fixed price",
            "status_label": "In progress",
            "estimator": "E",
            "revised_contract": D("100.00"),
            "deposit_invoiced": D("0.00"),
            "deposit_received": D("0.00"),
            "billed_to_date": D("-1234.50"),
            "fuel_surcharge_billed": D("0.00"),
            "collected_to_date": D("0.00"),
            "other_credits_applied": D("0.00"),
            "open_ar": D("0.00"),
            "remaining_to_bill": None,
            "days_since_activity": None,
        },
        notes={
            "remaining_to_bill": "Not shown: Pool has no contract.",
            "days_since_activity": "No activity",
        },
    )
    tie = TieRow("2026-09", D("1.00"), D("2.00"), D("3.00"), D("0.00"), D("0.00"), D("0.00"))
    report = export.BoardReport(
        "Co", date(2026, 10, 6), "All jobs", (row,), row, row, (tie,), "Ties.", None
    )
    xlsx = load_workbook(io.BytesIO(export.board_xlsx(report)))["Sold Jobs Board"]
    header, rows = _sheet_rows(xlsx)
    cells = dict(zip(header, rows[0], strict=True))
    assert (
        cells["Billed to date"] == D("-1234.50")
        and cells["Remaining to bill"] == "Not shown: Pool has no contract."
    )
    pdf = export.board_pdf(report).decode("latin-1")
    assert "A job" in pdf and "1,234.50" in pdf and "Not shown: Pool has no contract." in pdf
