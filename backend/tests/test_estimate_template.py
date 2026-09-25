"""F06: the estimate template parser on the Rye Beach fixtures and on test copies.
Every number below is in the brief (`current-feature.md`, Acceptance criteria)."""

from decimal import Decimal

import pytest

from app.domain.estimates.names import reads_as_unit_price, suggests_change_order
from app.integrations.estimate_template import TemplateError, parse_template
from tests.estimate_helpers import (
    EIGHTY,
    ELM,
    TEMPLATE,
    TURLEY,
    build_csv,
    build_workbook,
    by_entity,
    fixture_rows,
    parse_bytes,
)

D = Decimal


def _lines_by_order(payload: dict) -> dict[int, list[dict]]:
    out: dict[int, list[dict]] = {}
    for ln in payload["lines"]:
        out.setdefault(ln["order"], []).append(ln)
    return out


# --- the 67 Elm Street fixture (EST6115758) -------------------------------------------------------
def test_67_elm_parses_to_the_three_raw_items_with_the_replan_numbers() -> None:
    items, rejected = parse_bytes(ELM.read_bytes())
    assert rejected == []
    got = by_entity(items)
    assert set(got) == {
        ("estimate", "EST6115758"),
        ("work_areas", "EST6115758"),
        ("cost_lines", "EST6115758"),
    }
    est = got[("estimate", "EST6115758")]
    assert (est["status"], est["status_norm"], est["price"]) == ("Sold", "sold", D("475129.68"))
    assert isinstance(est["price"], Decimal)
    was = got[("work_areas", "EST6115758")]["work_areas"]
    assert len(was) == 21 and [w["order"] for w in was] == list(range(1, 22))
    omitted = [w["order"] for w in was if not w["kept"]]
    assert omitted == [17] and next(w for w in was if w["order"] == 17)["price"] == D("15000.00")
    kept = [w for w in was if w["kept"]]
    assert sum((w["price"] for w in kept), D(0)) == D("475129.68")
    assert next(w for w in was if w["order"] == 20)["price"] == D("0.00")
    assert all(isinstance(w["price"], Decimal) for w in was)
    lines = _lines_by_order(got[("cost_lines", "EST6115758")])
    # 60 rows on the sheet; work area 1 carries two 250 lines, summed at parse into one.
    assert sum(len(v) for v in lines.values()) == 59
    assert sorted(ln["cost_code"] for ln in lines[1]).count("250") == 1
    assert 17 not in lines and 20 not in lines
    assert [ln["cost_code"] for ln in lines[5]] == ["290"]  # the placeholder
    assert lines[5][0]["amount"] == D("46381.35")
    total = sum((ln["amount"] for v in lines.values() for ln in v), D(0))
    assert total == D("315832.69")
    assert all(isinstance(ln["amount"], Decimal) for v in lines.values() for ln in v)


def test_67_elm_name_rules_flag_18_to_21_and_not_the_concrete_rows() -> None:
    items, _ = parse_bytes(ELM.read_bytes())
    was = by_entity(items)[("work_areas", "EST6115758")]["work_areas"]
    flagged = [w["order"] for w in was if suggests_change_order(w["name"])]
    assert flagged == [18, 19, 20, 21]
    assert [w["order"] for w in was if reads_as_unit_price(w["name"])] == [18]
    names = {w["order"]: w["name"] for w in was}
    assert names[13].startswith("Concrete") and names[14].startswith("Concrete")


# --- the Turley fixture (EST6120638) --------------------------------------------------------------
def test_turley_parses_to_32_work_areas_seven_omitted_and_20_lines_on_190() -> None:
    items, rejected = parse_bytes(TURLEY.read_bytes())
    assert rejected == []
    got = by_entity(items)
    est = got[("estimate", "EST6120638")]
    assert (est["status_norm"], est["price"]) == ("sold", D("366889.80"))
    was = got[("work_areas", "EST6120638")]["work_areas"]
    assert len(was) == 32
    assert [w["order"] for w in was if not w["kept"]] == [6, 7, 10, 20, 21, 22, 23]
    assert sum((w["price"] for w in was if w["kept"]), D(0)) == D("366889.80")
    assert sum((w["price"] for w in was if not w["kept"]), D(0)) == D("186727.99")
    assert [w["order"] for w in was if suggests_change_order(w["name"])] == list(range(23, 33))
    assert [w["order"] for w in was if w["price"] == 0] == [26, 27, 28, 29, 30]
    lines = got[("cost_lines", "EST6120638")]["lines"]
    assert len(lines) == 20 and {ln["cost_code"] for ln in lines} == {"190"}
    assert sum((ln["amount"] for ln in lines), D(0)) == D("212639.82")
    assert {ln["order"] for ln in lines}.isdisjoint({6, 7, 10, 20, 21, 22, 23, 26, 27, 28, 29, 30})


# --- the 80-row Estimates workbook ----------------------------------------------------------------
@pytest.mark.skipif(not EIGHTY.exists(), reason="owner supplies estimates_2026-09-17.xlsx")
def test_80_rows_load_79_estimates_and_the_mijal_row_raises_est_no_id() -> None:
    items, rejected = parse_bytes(EIGHTY.read_bytes())
    assert [r.code for r in rejected] == ["EST_NO_ID"]
    assert rejected[0].detail["cells"]["price"] == "8548.03"
    assert "no estimate id" in rejected[0].message
    ests = [i.payload for i in items if i.entity_type == "estimate"]
    assert len(ests) == 79 and len(items) == 79
    by = {}
    for e in ests:
        key = (e["estimator"], e["status_norm"])
        n, total = by.get(key, (0, D(0)))
        by[key] = (n + 1, total + e["price"])
    assert by[("William Hess", "pending")] == (23, D("493619.70"))
    assert by[("William Hess", "sold")] == (10, D("429962.74"))
    assert by[("William Hess", "lost")] == (1, D("13078.03"))
    assert by[("Stephanie Sanford", "pending")] == (29, D("921867.83"))
    # Sanford's sold total on the report, 486,436.83, includes the id-less Mijal row.
    assert by[("Stephanie Sanford", "sold")] == (5, D("486436.83") - D("8548.03"))
    assert by[("Stephanie Sanford", "lost")] == (11, D("190798.89"))


# --- money never touches a float ------------------------------------------------------------------
def test_float_cells_are_read_exactly_and_rounded_half_up_to_the_cent() -> None:
    content = build_workbook(
        estimates=[
            {"estimate_id": "EST1", "name": "a", "status": "Lost", "price": 190798.88999999998},
            {"estimate_id": "EST2", "name": "b", "status": "Pending", "price": 1599103.5499999998},
            {"estimate_id": "EST3", "name": "c", "status": "Pending", "price": "1,234.505"},
            {"estimate_id": "EST4", "name": "d", "status": "Pending", "price": 0},
        ]
    )
    items, rejected = parse_bytes(content)
    assert rejected == []
    prices = {i.external_id: i.payload["price"] for i in items}
    assert prices == {
        "EST1": D("190798.89"),
        "EST2": D("1599103.55"),
        "EST3": D("1234.51"),
        "EST4": D("0.00"),
    }
    assert all(isinstance(p, Decimal) for p in prices.values())


def test_same_code_lines_under_one_work_area_are_summed_at_parse() -> None:
    content = build_workbook(
        estimates=[{"estimate_id": "EST1", "name": "a", "status": "Sold", "price": "10.00"}],
        work_areas=[
            {"estimate_id": "EST1", "order": 1, "kept": "Y", "name": "One", "price": "10.00"}
        ],
        costs=[
            {
                "estimate_id": "EST1",
                "order": 1,
                "cost_code": "130",
                "hours": 1,
                "amount": "2.50",
                "notes": "mulch",
            },
            {
                "estimate_id": "EST1",
                "order": 1,
                "cost_code": "130",
                "hours": "1.5",
                "amount": "3.25",
                "notes": "loam",
            },
            {"estimate_id": "EST1", "order": 1, "cost_code": "110", "amount": "1.00"},
        ],
    )
    items, rejected = parse_bytes(content)
    assert rejected == []
    lines = by_entity(items)[("cost_lines", "EST1")]["lines"]
    assert [(ln["cost_code"], ln["hours"], ln["amount"], ln["notes"]) for ln in lines] == [
        ("110", None, D("1.00"), None),
        ("130", D("2.50"), D("5.75"), "mulch; loam"),
    ]


# --- rows that do not load, one at a time ---------------------------------------------------------
def _one_estimate(**over) -> dict:
    row = {"estimate_id": "EST1", "name": "a", "status": "Sold", "price": "10.00"}
    row.update(over)
    return row


def test_a_row_without_an_estimate_id_is_rejected_with_its_cells_and_the_rest_loads() -> None:
    content = build_workbook(
        estimates=[
            _one_estimate(),
            {
                "estimator": "Pat",
                "client": "X",
                "jobsite": "Y",
                "name": "no id",
                "status": "Sold",
                "price": "8548.03",
            },
            _one_estimate(estimate_id="EST2"),
        ]
    )
    items, rejected = parse_bytes(content)
    assert [i.external_id for i in items] == ["EST1", "EST2"]
    (r,) = rejected
    assert (r.code, r.row_number) == ("EST_NO_ID", 3)
    assert r.detail["cells"] == {
        "estimate_id": "",
        "estimator": "Pat",
        "client": "X",
        "jobsite": "Y",
        "name": "no id",
        "status": "Sold",
        "price": "8548.03",
        "estimate_date": "",
    }
    assert r.message.startswith('Row 3 on "Estimates" has no estimate id')


def test_a_malformed_price_rejects_that_row_only() -> None:
    content = build_workbook(
        estimates=[
            _one_estimate(),
            _one_estimate(estimate_id="EST2", price="ten"),
            _one_estimate(estimate_id="EST3"),
        ]
    )
    items, rejected = parse_bytes(content)
    assert [i.external_id for i in items] == ["EST1", "EST3"]
    assert len(rejected) == 1 and rejected[0].code is None
    assert rejected[0].message == 'Row 3 on "Estimates" was not loaded: not a number.'
    assert rejected[0].detail == {"sheet": "Estimates", "row": 3, "estimate_id": "EST2"}


def test_kept_other_than_y_or_n_rejects_the_row() -> None:
    content = build_workbook(
        estimates=[_one_estimate()],
        work_areas=[
            {"estimate_id": "EST1", "order": 1, "kept": "Y", "name": "One", "price": "5.00"},
            {"estimate_id": "EST1", "order": 2, "kept": "yes", "name": "Two", "price": "5.00"},
            {"estimate_id": "EST1", "order": 3, "kept": "n", "name": "Three", "price": "1.00"},
        ],
    )
    items, rejected = parse_bytes(content)
    was = by_entity(items)[("work_areas", "EST1")]["work_areas"]
    assert [(w["order"], w["kept"]) for w in was] == [(1, True), (3, False)]
    assert [r.message for r in rejected] == [
        'Row 3 on "Work areas" was not loaded: kept must be Y or N.'
    ]


def test_an_unknown_status_word_loads_with_no_normalized_status() -> None:
    items, rejected = parse_bytes(build_workbook(estimates=[_one_estimate(status="Won ")]))
    assert rejected == []
    assert (items[0].payload["status"], items[0].payload["status_norm"]) == ("Won", None)
    assert (
        parse_bytes(build_workbook(estimates=[_one_estimate(status="sold")]))[0][0].payload[
            "status_norm"
        ]
        == "sold"
    )


def test_a_cost_line_for_a_work_area_not_on_the_sheet_in_the_same_file_is_rejected() -> None:
    content = build_workbook(
        estimates=[_one_estimate()],
        work_areas=[
            {"estimate_id": "EST1", "order": 1, "kept": "Y", "name": "One", "price": "10.00"}
        ],
        costs=[
            {"estimate_id": "EST1", "order": 1, "cost_code": "130", "amount": "1.00"},
            {"estimate_id": "EST1", "order": 9, "cost_code": "130", "amount": "1.00"},
        ],
    )
    items, rejected = parse_bytes(content)
    assert len(by_entity(items)[("cost_lines", "EST1")]["lines"]) == 1
    assert rejected[0].message == (
        'Row 3 on "Estimate costs" was not loaded: work area #9 of EST1 is not on '
        '"Work areas" in this file.'
    )


def test_a_duplicate_order_and_a_duplicate_estimate_row_are_rejected() -> None:
    content = build_workbook(
        estimates=[_one_estimate(), _one_estimate(name="again")],
        work_areas=[
            {"estimate_id": "EST1", "order": 1, "kept": "Y", "name": "One", "price": "5.00"},
            {"estimate_id": "EST1", "order": 1, "kept": "Y", "name": "One again", "price": "5.00"},
        ],
    )
    items, rejected = parse_bytes(content)
    assert sorted(r.message for r in rejected) == [
        'Row 3 on "Estimates" was not loaded: EST1 appears twice on this sheet.',
        'Row 3 on "Work areas" was not loaded: work area #1 of EST1 appears twice.',
    ]
    assert by_entity(items)[("estimate", "EST1")]["name"] == "a"


def test_a_costs_only_file_keeps_its_lines_for_the_normalizer() -> None:
    """Without Work areas rows in the same file the parser cannot check the order;
    the normalizer checks it against the spine."""
    items, rejected = parse_bytes(
        build_csv(
            "Estimate costs",
            [{"estimate_id": "EST1", "order": 4, "cost_code": "130", "amount": "1.00"}],
        )
    )
    assert rejected == [] and [(i.entity_type, i.external_id) for i in items] == [
        ("cost_lines", "EST1")
    ]


# --- the CSV form and the workbook's shape --------------------------------------------------------
def test_each_csv_is_recognised_by_its_header() -> None:
    est = build_csv("Estimates", [_one_estimate(estimate_date="2026-09-17")])
    was = build_csv(
        "Work areas",
        [{"estimate_id": "EST1", "order": 1, "kept": "N", "name": "One", "price": "10.00"}],
    )
    cost = build_csv(
        "Estimate costs",
        [{"estimate_id": "EST1", "order": 1, "cost_code": "130", "amount": "1.00"}],
    )
    assert [(i.entity_type) for i in parse_bytes(est)[0]] == ["estimate"]
    assert parse_bytes(est)[0][0].payload["estimate_date"] == "2026-09-17"
    assert [(i.entity_type) for i in parse_bytes(was)[0]] == ["work_areas"]
    assert [(i.entity_type) for i in parse_bytes(cost)[0]] == ["cost_lines"]


def test_a_workbook_without_a_template_sheet_or_a_required_column_fails_the_parse() -> None:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(TemplateError):
        list(parse_template(io.BytesIO(buf.getvalue())))
    wb = Workbook()
    wb.active.title = "Estimates"
    wb.active.append(["estimate_id", "name", "price"])  # no status
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(TemplateError):
        list(parse_template(io.BytesIO(buf.getvalue())))
    with pytest.raises(TemplateError):
        list(parse_template(io.BytesIO(b"a,b\n1,2\n")))


def test_extra_columns_and_blank_rows_are_ignored_and_headers_match_loosely() -> None:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = " estimates "
    ws.append([" Estimate_ID ", "Name", "STATUS", "Price ", "Extra"])
    ws.append([None, None, None, None, None])
    ws.append(["EST9", "x", "Lost", 1.5, "ignored"])
    buf = io.BytesIO()
    wb.save(buf)
    items, rejected = parse_bytes(buf.getvalue())
    assert rejected == [] and items[0].payload["price"] == D("1.50")


def test_the_blank_template_parses_to_nothing() -> None:
    if not TEMPLATE.exists():
        pytest.skip("docs/templates/estimate_template.xlsx is generated at build")
    assert parse_bytes(TEMPLATE.read_bytes()) == ([], [])


def test_fixture_rows_helper_round_trips_the_turley_fixture() -> None:
    rows = fixture_rows(TURLEY)
    copy = build_workbook(sheets=rows)
    assert by_entity(parse_bytes(copy)[0]) == by_entity(parse_bytes(TURLEY.read_bytes())[0])
