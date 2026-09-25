"""Helpers shared by the F06 tests: the Rye Beach template fixtures, workbook and
CSV builders for test copies, the parser on bytes, uploading through the API, and
reading the estimate tables back."""

import csv
import io
import uuid
from decimal import Decimal
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import Engine, select

from app.core.db import tenant_session
from app.domain.config.audit import Actor
from app.domain.config.policy import set_policy
from app.domain.estimates.models import Estimate, EstimateCost, EstimateVersion, EstimateWorkArea
from app.domain.estimates.normalize import NORMALIZE  # registers the task for the test worker
from app.integrations.base import RawItem, RejectedItem
from app.integrations.estimate_template import (
    COST_COLUMNS,
    ESTIMATE_COLUMNS,
    WORK_AREA_COLUMNS,
    parse_template,
)
from tests.config_helpers import load_rye_beach_rules
from tests.conftest import CSRF, Seed

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "rye_beach" / "estimates"
ELM = FIXTURES / "estimate_upload_EST6115758.xlsx"
TURLEY = FIXTURES / "estimate_upload_EST6120638.xlsx"
EIGHTY = FIXTURES / "estimates_2026-09-17.xlsx"
TEMPLATE = Path(__file__).resolve().parents[2] / "docs" / "templates" / "estimate_template.xlsx"

# D-04: the WIP basis for Rye Beach, as slots (Labor Burden and Warranty undecided → out).
D04_BASIS = ["10", "30", "35", "40", "50", "60", "70", "90"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TASK_KIND = NORMALIZE

SHEETS = {
    "Estimates": ESTIMATE_COLUMNS,
    "Work areas": WORK_AREA_COLUMNS,
    "Estimate costs": COST_COLUMNS,
}


def fixture_rows(path: Path) -> dict[str, list[dict]]:
    """Every data sheet of a fixture as a list of dicts keyed by header (cells as
    the workbook holds them), so tests can build a changed copy."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[dict]] = {}
    try:
        for ws in wb.worksheets:
            if ws.title not in SHEETS:
                continue
            rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
            header = [str(h).strip() for h in rows[0]]
            out[ws.title] = [dict(zip(header, r, strict=False)) for r in rows[1:]]
    finally:
        wb.close()
    return out


def build_workbook(
    estimates: list[dict] | None = None,
    work_areas: list[dict] | None = None,
    costs: list[dict] | None = None,
    *,
    readme: bool = True,
    sheets: dict[str, list[dict]] | None = None,
) -> bytes:
    """A template workbook from row dicts (missing keys are blank cells)."""
    data = sheets or {}
    if estimates is not None:
        data["Estimates"] = estimates
    if work_areas is not None:
        data["Work areas"] = work_areas
    if costs is not None:
        data["Estimate costs"] = costs
    wb = Workbook()
    first = wb.active
    if readme:
        first.title = "Read me"
        first.append(["Estimate upload template (test copy)"])
    else:
        wb.remove(first)
    for title, columns in SHEETS.items():
        if title not in data:
            continue
        ws = wb.create_sheet(title)
        ws.append(list(columns))
        for row in data[title]:
            ws.append([row.get(c) for c in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(sheet: str, rows: list[dict]) -> bytes:
    columns = SHEETS[sheet]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for row in rows:
        w.writerow(["" if row.get(c) is None else row.get(c) for c in columns])
    return buf.getvalue().encode()


def parse_bytes(content: bytes) -> tuple[list[RawItem], list[RejectedItem]]:
    items = list(parse_template(io.BytesIO(content)))
    return (
        [i for i in items if isinstance(i, RawItem)],
        [i for i in items if isinstance(i, RejectedItem)],
    )


def by_entity(items: list[RawItem]) -> dict[tuple[str, str], dict]:
    return {(i.entity_type, i.external_id): i.payload for i in items}


def upload_template(client: TestClient, content: bytes, filename: str = "estimates.xlsx") -> dict:
    ctype = XLSX if filename.endswith(".xlsx") else "text/csv"
    r = client.post(
        "/api/imports",
        data={"source_kind": "estimate_template"},
        files={"file": (filename, io.BytesIO(content), ctype)},
        headers=CSRF,
    )
    assert r.status_code in (200, 201), r.text
    return r.json()


def configure_tenant(
    engine: Engine, seed: Seed, tenant_id: uuid.UUID, *, basis: bool = True
) -> None:
    """Rye Beach's divisions (LS 1, EX 2, GC 3, SNOW 4 from the rules fixture) and,
    unless ``basis`` is False, the D-04 WIP basis policy."""
    load_rye_beach_rules(engine, tenant_id)
    if basis:
        with tenant_session(engine, tenant_id) as s:
            set_policy(
                s,
                tenant_id,
                "wip_basis",
                D04_BASIS,
                decision_ref="D-04 (test)",
                actor=Actor(user_id=seed.users["rotate_me"].id),
            )


def estimate_rows(engine: Engine, tenant_id: uuid.UUID) -> dict[str, Estimate]:
    with tenant_session(engine, tenant_id) as s:
        rows = {e.external_id: e for e in s.execute(select(Estimate)).scalars()}
        for e in rows.values():
            s.expunge(e)
        return rows


def versions_of(engine: Engine, tenant_id: uuid.UUID, external_id: str) -> list[EstimateVersion]:
    with tenant_session(engine, tenant_id) as s:
        est = s.execute(select(Estimate).where(Estimate.external_id == external_id)).scalar_one()
        rows = list(
            s.execute(
                select(EstimateVersion)
                .where(EstimateVersion.estimate_id == est.id)
                .order_by(EstimateVersion.version_no)
            ).scalars()
        )
        for v in rows:
            s.expunge(v)
        return rows


def work_areas_of(engine: Engine, tenant_id: uuid.UUID, version_id: uuid.UUID) -> list[dict]:
    """The work areas of a version with their cost lines, as plain dicts."""
    with tenant_session(engine, tenant_id) as s:
        was = list(
            s.execute(
                select(EstimateWorkArea)
                .where(EstimateWorkArea.estimate_version_id == version_id)
                .order_by(EstimateWorkArea.order_no)
            ).scalars()
        )
        out = []
        for w in was:
            lines = list(
                s.execute(
                    select(EstimateCost)
                    .where(EstimateCost.estimate_work_area_id == w.id)
                    .order_by(EstimateCost.cost_code)
                ).scalars()
            )
            out.append(
                {
                    "order": w.order_no,
                    "name": w.name,
                    "kept": w.kept,
                    "co": w.change_order_suggested,
                    "price": w.price,
                    "lines": [
                        {
                            "code": ln.cost_code,
                            "hours": ln.hours,
                            "amount": ln.amount,
                            "category": ln.cost_category_id,
                            "division": ln.division_id,
                        }
                        for ln in lines
                    ],
                }
            )
        return out


def money(text: str) -> Decimal:
    return Decimal(text)
