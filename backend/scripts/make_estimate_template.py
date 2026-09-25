#!/usr/bin/env python
"""Write the blank estimate template (F06, D-32): ``docs/templates/estimate_template.xlsx``.

    cd backend && .venv/bin/python scripts/make_estimate_template.py

A *Read me* sheet with the column meanings and the three data sheets with their
header rows and nothing else. The Read me text lives here, in one place; the parser
(``app.integrations.estimate_template``) reads the sheets by name and the columns by
header, so the file this script writes is by construction what the parser accepts
(``tests/test_estimate_template.py`` parses it to nothing).
"""

import sys
from pathlib import Path

from openpyxl import Workbook

from app.integrations.estimate_template import COST_COLUMNS, ESTIMATE_COLUMNS, WORK_AREA_COLUMNS

OUT = Path(__file__).resolve().parents[2] / "docs" / "templates" / "estimate_template.xlsx"

README: tuple[str, ...] = (
    "Estimate upload template",
    "One workbook carries any number of estimates. Fill the three data sheets; leave this "
    "sheet as it is.",
    "",
    "Sheet 'Estimates' — one row per estimate.",
    "  estimate_id     required. Your estimating system's id (e.g. EST6115758). Rows are "
    "matched and updated by this id, never by name.",
    "  estimator, client, jobsite   optional text. The platform never links a client by name.",
    "  name            required. Estimate / project name.",
    "  status          required. Pending, Sold or Lost. Any other word is loaded without a "
    "normalized status and reported.",
    "  price           required. The estimate's total price: the control total. Kept work "
    "areas on 'Work areas' must sum to it.",
    "  estimate_date   optional. YYYY-MM-DD.",
    "",
    "Sheet 'Work areas' — one row per work area (the schedule-of-values line the customer "
    "sees and that billing refers to as #n).",
    "  estimate_id, order   required. order is the line number; it is the work area's "
    "identity across versions — never renumber.",
    "  kept            required. Y or N. N = omitted from the contract. An omitted work "
    "area, or one priced 0.00, needs no cost lines.",
    "  name            required. A name beginning 'CO:' or 'C/O' suggests a change order; "
    "a person confirms it.",
    "  price           required. notes optional.",
    "",
    "Sheet 'Estimate costs' — one or more rows per kept work area, each with exactly one "
    "cost code. Two rows with the same code under one work area are summed.",
    "  estimate_id, order   required; must match a row on 'Work areas'.",
    "  cost_code       required. Three digits: division digit + cost-type slot (e.g. 210 = "
    "EX Labor, 240 = EX Subcontractors). The code equals the ledger account without its "
    "leading 5.",
    "  hours           optional. Estimated labor (or machine) hours for the line.",
    "  amount          required. Estimated cost. A work area's cost is the sum of its "
    "lines; the estimate's cost by category is the sum over kept work areas.",
    "  notes           optional. What the line is made of.",
    "",
    "Money to the cent. A kept work area with no cost line leaves the estimate outside the "
    "WIP schedule until a person adds the lines (D-04).",
    "Re-uploading the same file changes nothing; a changed file creates a new version and "
    "the earlier one is kept.",
    "The file may also be uploaded as three .csv files, one per sheet, with the same "
    "column headers, in any order.",
)


def build() -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for line in README:
        readme.append([line])
    readme.column_dimensions["A"].width = 120
    for title, columns in (
        ("Estimates", ESTIMATE_COLUMNS),
        ("Work areas", WORK_AREA_COLUMNS),
        ("Estimate costs", COST_COLUMNS),
    ):
        ws = wb.create_sheet(title)
        ws.append(list(columns))
        for i, name in enumerate(columns, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
                14, len(name) + 4
            )
        ws.freeze_panes = "A2"
    return wb


def main(argv: list[str]) -> int:
    out = Path(argv[1]) if len(argv) > 1 else OUT
    out.parent.mkdir(parents=True, exist_ok=True)
    build().save(out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
